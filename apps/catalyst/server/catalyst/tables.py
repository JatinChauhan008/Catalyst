"""Extract every table of a converted document, with metadata about its origin.

For each table the extractor records where it comes from (document, page,
position, the title above it, its section, caption and units), how it was read
(PDF text layer or OCR) and whether it needs a human check. Tables that
continue on the next page are linked, or joined when asked.

`write_table_package` turns the records into files: a CSV per table, an Excel
workbook with an index sheet, a JSON manifest and an image of each table.
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from itertools import pairwise
from pathlib import Path

from docling_core.types.doc import (
    DocItem,
    DocItemLabel,
    DoclingDocument,
    NodeItem,
    PictureItem,
    TableItem,
    TextItem,
)

from catalyst.models import (
    TableExtraction,
    TableExtractionSettings,
    TableFileType,
    TableManifest,
    TableQuality,
    TableRecord,
    TableRegion,
    TableSource,
)

# Text right above a table that names it ("STATEMENT SHOWING DISTRICT-WISE ...").
_TITLE_MAX_CHARS = 200
_SENTENCE_MIN_WORDS = 8
# A label such as "[Table-13]" that the layout model marked as a heading.
_TABLE_LABEL = re.compile(r"^\[?\s*table\s*-?\s*\d+\s*\]?$", re.IGNORECASE)
_TITLE_LOOKBACK = 4
_ABOVE_TOLERANCE = 6.0  # points

_UNITS = re.compile(
    r"\((?:in|rs\.?)\s[^)]*\)"
    r"|\b(?:rs\.?|amount|amt\.?)\s+in\s+[a-z]+"
    r"|\bin\s+(?:lakhs?|crores?|thousands?|millions?|billions?|rupees|mm|mt|hects?|hectares|tonnes|percent)\b",
    re.IGNORECASE,
)

# A continued table starts high on the next page, after its first part ended low.
_CONTINUED_END = 0.6
_CONTINUED_START = 0.4

_EMPTY_SHARE_WARN = 0.6

# Formats whose "pages" are sheets and whose boxes are cell positions.
_SPREADSHEETS = frozenset({"xlsx", "xls", "ods", "csv", "iwork_numbers"})


@dataclass
class ExtractionContext:
    """What the extractor needs to know about the conversion."""

    source: TableSource
    is_pdf: bool
    is_image: bool
    text_layer_ignored: bool = False
    pdf_bytes: bytes | None = None
    ocr_engine: str | None = None
    table_text: str | None = None


@dataclass
class _Entry:
    record: TableRecord
    items: list[TableItem] = field(default_factory=list)
    # A title or a caption above the table: it starts a table of its own.
    labelled_above: bool = False


def extract_tables(
    document: DoclingDocument,
    context: ExtractionContext,
    settings: TableExtractionSettings,
) -> tuple[list[TableRecord], dict[str, list[TableItem]]]:
    """Records for all tables in reading order, and the Docling items behind each."""
    items: list[NodeItem] = [
        item for item, _ in document.iterate_items(with_groups=False)
    ]
    text_layer = (
        _TextLayer(context.pdf_bytes) if context.is_pdf and context.pdf_bytes else None
    )

    entries: list[_Entry] = []
    per_page: dict[int, int] = {}
    for position, item in enumerate(items):
        if not isinstance(item, TableItem):
            continue
        rows = [[cell.text for cell in row] for row in item.data.grid]
        units = "cell" if context.source.detected_format in _SPREADSHEETS else "pt"
        regions = _regions(item, document, units)
        page = regions[0].page if regions else None
        if page is not None:
            per_page[page] = per_page.get(page, 0) + 1
            table_id = f"{context.source.sha256[:12]}-p{page:03d}-t{per_page[page]:02d}"
        else:
            table_id = f"{context.source.sha256[:12]}-t{len(entries) + 1:03d}"

        title = _title(items, position, item, document)
        caption = item.caption_text(document) or None
        header_rows = _header_rows(item)
        cells = [text for row in rows for text in row]
        empty_share = (
            (sum(1 for text in cells if not text.strip()) / len(cells))
            if cells
            else 1.0
        )
        record = TableRecord(
            table_id=table_id,
            index=len(entries) + 1,
            source=context.source,
            page=page,
            regions=regions,
            caption=caption,
            title=title,
            section=_section(items, position, title),
            units=_units(
                items, position, item, document, [title, caption, *rows[:header_rows]]
            ),
            footnotes=[ref.resolve(document).text for ref in item.footnotes],
            num_rows=len(rows),
            num_cols=item.data.num_cols,
            header_rows=header_rows,
            empty_cells_share=round(empty_share, 3),
            extraction=TableExtraction(
                method=_method(context, regions, text_layer),
                ocr_engine=context.ocr_engine,
                table_text=context.table_text,
            ),
            quality=_quality(len(rows), item.data.num_cols, empty_share),
            preview=rows[: settings.preview_rows],
            rows=rows,
        )
        entries.append(
            _Entry(record, [item], title is not None or _caption_above(item, document))
        )

    _link_continued(entries)
    if settings.join_continued:
        entries = _join_continued(entries, settings)
    return [entry.record for entry in entries], {
        e.record.table_id: e.items for e in entries
    }


# ------------------------------------------------------------------ metadata
def _regions(item: DocItem, document: DoclingDocument, units: str) -> list[TableRegion]:
    regions = []
    for prov in item.prov:
        page = document.pages.get(prov.page_no)
        if page is None:
            continue
        width, height = page.size.width, page.size.height
        box = prov.bbox.to_top_left_origin(page_height=height)
        regions.append(
            TableRegion(
                page=prov.page_no,
                bbox=(
                    round(box.l, 2),
                    round(box.t, 2),
                    round(box.r, 2),
                    round(box.b, 2),
                ),
                bbox_normalized=(
                    round(box.l / width, 4),
                    round(box.t / height, 4),
                    round(box.r / width, 4),
                    round(box.b / height, 4),
                ),
                page_width=width,
                page_height=height,
                units=units,
            )
        )
    return regions


def _top(item: NodeItem, document: DoclingDocument) -> tuple[int, float, float] | None:
    """(page, top, bottom) of an item's first region, top-left origin."""
    if not isinstance(item, TextItem | TableItem | PictureItem) or not item.prov:
        return None
    prov = item.prov[0]
    page = document.pages.get(prov.page_no)
    if page is None:
        return None
    box = prov.bbox.to_top_left_origin(page_height=page.size.height)
    return prov.page_no, box.t, box.b


def _title(
    items: list[NodeItem], position: int, table: TableItem, document: DoclingDocument
) -> str | None:
    """The short text right above the table (on its page, when positions are known)."""
    where = _top(table, document)
    captions = {ref.cref for ref in table.captions}
    for item in reversed(items[max(0, position - _TITLE_LOOKBACK) : position]):
        if isinstance(item, TableItem | PictureItem):
            return None
        if not isinstance(item, TextItem) or item.self_ref in captions:
            continue
        if where is not None:
            found = _top(item, document)
            if (
                found is None
                or found[0] != where[0]
                or found[2] > where[1] + _ABOVE_TOLERANCE
            ):
                continue
        elif item.label in (DocItemLabel.SECTION_HEADER, DocItemLabel.TITLE):
            # Without positions (Markdown, Word), a heading is the section, not a title.
            return None
        text = item.text.strip()
        if (
            item.label == DocItemLabel.TEXT
            and text.endswith(".")
            and len(text.split()) >= _SENTENCE_MIN_WORDS
        ):
            return None  # a sentence of the body text, not a title
        if (
            text
            and len(text) <= _TITLE_MAX_CHARS
            and item.label != DocItemLabel.LIST_ITEM
        ):
            return text
    return None


def _caption_above(table: TableItem, document: DoclingDocument) -> bool:
    where = _top(table, document)
    if where is None:
        return False
    for ref in table.captions:
        found = _top(ref.resolve(document), document)
        if found and found[0] == where[0] and found[2] <= where[1] + _ABOVE_TOLERANCE:
            return True
    return False


def _section(items: list[NodeItem], position: int, title: str | None) -> str | None:
    """The nearest heading before the table that is not its title, else the title heading."""
    title_heading = None
    for item in reversed(items[:position]):
        if isinstance(item, TextItem) and item.label in (
            DocItemLabel.SECTION_HEADER,
            DocItemLabel.TITLE,
        ):
            text = item.text.strip()
            if _TABLE_LABEL.match(text):
                continue
            if text != title:
                return text
            title_heading = title_heading or text
    return title_heading


def _units(
    items: list[NodeItem],
    position: int,
    table: TableItem,
    document: DoclingDocument,
    nearby: list,
) -> str | None:
    """A units note ("Rs in lakh", "(IN MM)") in the title, caption, header or just above."""
    texts = [
        " ".join(part) if isinstance(part, list) else part for part in nearby if part
    ]
    where = _top(table, document)
    for item in reversed(items[max(0, position - _TITLE_LOOKBACK) : position]):
        if isinstance(item, TableItem | PictureItem):
            break
        if not isinstance(item, TextItem):
            continue
        found = _top(item, document)
        if where is None or (
            found and found[0] == where[0] and found[2] <= where[1] + _ABOVE_TOLERANCE
        ):
            texts.append(item.text)
    for text in texts:
        match = _UNITS.search(text)
        if match:
            return match.group(0).strip()
    return None


def _header_rows(table: TableItem) -> int:
    count = 0
    for row in table.data.grid:
        filled = [cell for cell in row if cell.text.strip()]
        if filled and all(cell.column_header for cell in filled):
            count += 1
        else:
            break
    return count


def _quality(rows: int, cols: int, empty_share: float) -> TableQuality:
    reasons = []
    if rows == 0 or cols == 0:
        reasons.append("no_cells")
    elif cols == 1 and rows >= 3:
        reasons.append("one_column: the columns may be merged")
    if rows and empty_share > _EMPTY_SHARE_WARN:
        reasons.append(
            f"mostly_empty: {round(empty_share * 100)}% of the cells are empty"
        )
    return TableQuality(status="warn" if reasons else "ok", reasons=reasons)


class _TextLayer:
    """Counts the PDF text inside a region, to tell text-layer tables from scans."""

    def __init__(self, pdf_bytes: bytes) -> None:
        import pypdfium2

        self._pdf = pypdfium2.PdfDocument(io.BytesIO(pdf_bytes))

    def chars(self, region: TableRegion) -> int:
        page = self._pdf[region.page - 1]
        left, top, right, bottom = region.bbox
        height = page.get_height()
        text = page.get_textpage().get_text_bounded(
            left, height - bottom, right, height - top
        )
        return len(text.strip())


def _method(
    context: ExtractionContext,
    regions: list[TableRegion],
    text_layer: _TextLayer | None,
) -> str:
    if context.is_image or context.text_layer_ignored:
        return "ocr"
    if not context.is_pdf:
        return "structured"
    if text_layer is None or not regions:
        return "ocr"
    with_text = [text_layer.chars(region) >= 10 for region in regions]
    if all(with_text):
        return "text_layer"
    return "mixed" if any(with_text) else "ocr"


# ------------------------------------------------------- tables across pages
def _header(record: TableRecord) -> list[list[str]]:
    return record.rows[: record.header_rows]


def _continues(first: TableRecord, second: TableRecord) -> bool:
    if not first.regions or not second.regions or first.num_cols != second.num_cols:
        return False
    end, start = first.regions[-1], second.regions[0]
    if start.page != end.page + 1:
        return False
    if (
        end.bbox_normalized[3] < _CONTINUED_END
        or start.bbox_normalized[1] > _CONTINUED_START
    ):
        return False
    return second.header_rows == 0 or _header(second) == _header(first)


def _link_continued(entries: list[_Entry]) -> None:
    for before, after in pairwise(entries):
        if not after.labelled_above and _continues(before.record, after.record):
            before.record.continues_to = after.record.table_id
            after.record.continues_from = before.record.table_id


def _join_continued(
    entries: list[_Entry], settings: TableExtractionSettings
) -> list[_Entry]:
    joined: list[_Entry] = []
    for entry in entries:
        previous = joined[-1] if joined else None
        last_part = (
            (previous.record.joined_from or [previous.record.table_id])[-1]
            if previous
            else None
        )
        if previous is not None and entry.record.continues_from == last_part:
            first, part = previous.record, entry.record
            rows = (
                part.rows[part.header_rows :]
                if _header(part) == _header(first)
                else part.rows
            )
            first.rows = first.rows + rows
            first.num_rows = len(first.rows)
            first.regions = first.regions + part.regions
            first.footnotes = first.footnotes + part.footnotes
            first.joined_from = (first.joined_from or [first.table_id]) + [
                part.table_id
            ]
            first.continues_to = part.continues_to
            first.preview = first.rows[: settings.preview_rows]
            if first.extraction.method != part.extraction.method:
                first.extraction.method = "mixed"
            previous.items.extend(entry.items)
            continue
        joined.append(entry)
    for index, entry in enumerate(joined, start=1):
        entry.record.index = index
    # Links to a part that was joined now point to the table that holds it.
    holder = {
        part: entry.record.table_id
        for entry in joined
        for part in entry.record.joined_from
    }
    for entry in joined:
        record = entry.record
        if record.continues_to:
            record.continues_to = holder.get(record.continues_to, record.continues_to)
        if record.continues_from:
            record.continues_from = holder.get(
                record.continues_from, record.continues_from
            )
    return joined


# --------------------------------------------------------------------- files
def write_table_package(
    records: list[TableRecord],
    items: dict[str, list[TableItem]],
    document: DoclingDocument,
    target: Path,
    settings: TableExtractionSettings,
    url_for: Callable[[Path], str],
) -> dict[str, Path]:
    """Write the package files. Fills `record.files`; returns the package-level files."""
    target.mkdir(parents=True, exist_ok=True)
    package: dict[str, Path] = {}

    for record in records:
        if TableFileType.CSV in settings.file_types:
            path = target / f"{record.table_id}.csv"
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                csv.writer(stream).writerows(record.rows)
            record.files["csv"] = url_for(path)
        if settings.include_images:
            images = [
                image
                for table in items[record.table_id]
                for index in range(len(table.prov))
                if (image := table.get_image(document, prov_index=index)) is not None
            ]
            for number, image in enumerate(images, start=1):
                name = (
                    f"{record.table_id}.png"
                    if number == 1
                    else f"{record.table_id}-{number}.png"
                )
                image.save(target / name)
                record.files["image" if number == 1 else f"image_{number}"] = url_for(
                    target / name
                )

    if TableFileType.XLSX in settings.file_types and records:
        package["tables_xlsx"] = target / "tables.xlsx"
        _write_workbook(records, items, package["tables_xlsx"])
        for record in records:
            record.files["xlsx"] = url_for(package["tables_xlsx"])
    if TableFileType.JSON in settings.file_types:
        package["tables_json"] = target / "tables.json"
        manifest = (
            TableManifest(source=records[0].source, tables=records) if records else None
        )
        if manifest is not None:
            package["tables_json"].write_text(
                manifest.model_dump_json(indent=2), encoding="utf-8"
            )
        else:
            package.pop("tables_json")
    return package


_INDEX_COLUMNS = [
    "Sheet", "Table ID", "Page(s)", "Title", "Caption", "Section", "Units",
    "Rows", "Columns", "Method", "Quality", "Notes", "Continues from", "Continues to",
]  # fmt: skip


def _write_workbook(
    records: list[TableRecord], items: dict[str, list[TableItem]], path: Path
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    workbook = Workbook()
    index = workbook.active
    index.title = "Index"
    index.append(_INDEX_COLUMNS)
    for cell in index[1]:
        cell.font = Font(bold=True)

    for record in records:
        pages = sorted({region.page for region in record.regions})
        sheet_name = f"T{record.index:02d}" + (f" p{pages[0]}" if pages else "")
        sheet = workbook.create_sheet(sheet_name[:31])
        for row in record.rows:
            sheet.append(row)
        for row in sheet.iter_rows(min_row=1, max_row=record.header_rows):
            for cell in row:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if not record.joined_from:
            _merge_spans(sheet, items[record.table_id][0])
        index.append(
            [
                sheet.title,
                record.table_id,
                ", ".join(map(str, pages)),
                record.title,
                record.caption,
                record.section,
                record.units,
                record.num_rows,
                record.num_cols,
                record.extraction.method,
                record.quality.status,
                "; ".join(record.quality.reasons),
                record.continues_from,
                record.continues_to,
            ]
        )
    workbook.save(path)


def _merge_spans(sheet, table: TableItem) -> None:
    """Merge the cells that span several rows or columns, as in the source."""
    seen = set()
    for cell in table.data.table_cells:
        if cell.row_span <= 1 and cell.col_span <= 1:
            continue
        key = (cell.start_row_offset_idx, cell.start_col_offset_idx)
        if key in seen:
            continue
        seen.add(key)
        sheet.merge_cells(
            start_row=cell.start_row_offset_idx + 1,
            start_column=cell.start_col_offset_idx + 1,
            end_row=cell.end_row_offset_idx,
            end_column=cell.end_col_offset_idx,
        )
