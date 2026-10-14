"""Table extraction: every table, with metadata that says where it comes from."""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path

import pypdfium2
from docling_core.types.doc import (
    BoundingBox,
    CoordOrigin,
    DoclingDocument,
    ProvenanceItem,
    Size,
    TableCell,
    TableData,
)
from extracta.models import (
    ConvertRequest,
    TableExtractionSettings,
    TableManifest,
    TableSource,
)
from extracta.tables import ExtractionContext, extract_tables
from openpyxl import load_workbook

TABLE_PAGE = (
    Path(__file__).resolve().parents[3]
    / "tests"
    / "data"
    / "pdf"
    / "sources"
    / "2305.03393v1-pg9.pdf"
)
TABLES_ON = {"export": {"to_formats": ["md"]}, "tables": {"enabled": True}}


def _convert(service, name: str, data: bytes, request: dict):
    return service.convert(name, data, ConvertRequest.model_validate(request))


def _file(service, envelope, url: str) -> Path:
    path = service.job_file(envelope.job_id, url.split("/files/", 1)[1])
    assert path is not None, url
    return path


def test_pdf_table_comes_with_its_origin_and_files(service) -> None:
    envelope = _convert(service, TABLE_PAGE.name, TABLE_PAGE.read_bytes(), TABLES_ON)
    assert envelope.status == "success", envelope.errors
    [table] = envelope.tables

    assert table.table_id.startswith(envelope.source.sha256[:12])
    assert table.page == 1
    [region] = table.regions
    left, top, right, bottom = region.bbox
    assert (
        0 <= left < right <= region.page_width
        and 0 <= top < bottom <= region.page_height
    )
    assert table.caption and table.caption.startswith("Table 1")
    assert table.extraction.method == "text_layer"
    assert table.header_rows >= 1 and table.num_cols >= 6
    assert table.quality.status == "ok"

    rows = list(
        csv.reader(
            io.StringIO(
                _file(service, envelope, table.files["csv"]).read_text("utf-8-sig")
            )
        )
    )
    assert len(rows) == table.num_rows
    assert _file(service, envelope, table.files["image"]).stat().st_size > 0

    workbook = load_workbook(_file(service, envelope, table.files["xlsx"]))
    assert workbook.sheetnames[0] == "Index" and len(workbook.sheetnames) == 2
    assert workbook["Index"]["B2"].value == table.table_id

    manifest_output = next(o for o in envelope.outputs if o.format == "tables_json")
    manifest = TableManifest.model_validate(json.loads(manifest_output.content))
    assert manifest.tables[0].rows == rows


def test_table_ids_are_stable_across_runs(service) -> None:
    data = TABLE_PAGE.read_bytes()
    first = _convert(service, TABLE_PAGE.name, data, TABLES_ON)
    second = _convert(service, "renamed.pdf", data, TABLES_ON)
    assert [t.table_id for t in first.tables] == [t.table_id for t in second.tables]


def test_scanned_table_has_the_same_shape_as_the_digital_one(service) -> None:
    digital = _convert(service, TABLE_PAGE.name, TABLE_PAGE.read_bytes(), TABLES_ON)
    scan = io.BytesIO()
    pypdfium2.PdfDocument(TABLE_PAGE)[0].render(scale=2.0).to_pil().save(
        scan, format="PNG"
    )
    request = {
        **TABLES_ON,
        "formats": {
            "image": {
                "pipeline_options": {"ocr_options": {"__variant__": "RapidOcrOptions"}}
            }
        },
    }
    scanned = _convert(service, "scan.png", scan.getvalue(), request)
    assert scanned.status == "success", scanned.errors
    assert len(scanned.tables) == len(digital.tables) == 1
    assert scanned.tables[0].extraction.method == "ocr"
    assert scanned.tables[0].num_cols == digital.tables[0].num_cols


# --------------------------------------------------- tables across two pages
H = 800.0


def _table(
    doc: DoclingDocument,
    page: int,
    top: float,
    bottom: float,
    rows: list[list[str]],
    header: bool,
):
    cells = [
        TableCell(
            text=text,
            start_row_offset_idx=r,
            end_row_offset_idx=r + 1,
            start_col_offset_idx=c,
            end_col_offset_idx=c + 1,
            column_header=header and r == 0,
        )
        for r, row in enumerate(rows)
        for c, text in enumerate(row)
    ]
    box = BoundingBox(l=50, t=top, r=550, b=bottom, coord_origin=CoordOrigin.TOPLEFT)
    doc.add_table(
        data=TableData(num_rows=len(rows), num_cols=len(rows[0]), table_cells=cells),
        prov=ProvenanceItem(
            page_no=page, bbox=box.to_bottom_left_origin(H), charspan=(0, 0)
        ),
    )


def _continued_report() -> DoclingDocument:
    doc = DoclingDocument(name="report")
    for page in (1, 2):
        doc.add_page(page_no=page, size=Size(width=600, height=H))
    header = ["Sl", "District", "Amount"]
    _table(
        doc,
        1,
        300,
        760,
        [header, ["1", "Angul", "605.66"], ["2", "Balasore", "219.70"]],
        header=True,
    )
    _table(doc, 2, 60, 400, [header, ["3", "Bargarh", "291.69"]], header=True)
    return doc


def _context() -> ExtractionContext:
    source = TableSource(
        filename="report.pdf", sha256="ab" * 32, detected_format="docx", page_count=2
    )
    return ExtractionContext(source=source, is_pdf=False, is_image=False)


def test_a_table_continued_on_the_next_page_is_linked() -> None:
    records, _ = extract_tables(
        _continued_report(), _context(), TableExtractionSettings(enabled=True)
    )
    first, second = records
    assert first.continues_to == second.table_id
    assert second.continues_from == first.table_id


def test_joined_table_keeps_one_header_and_both_regions() -> None:
    settings = TableExtractionSettings(enabled=True, join_continued=True)
    [table], _ = extract_tables(_continued_report(), _context(), settings)
    assert [row[1] for row in table.rows] == [
        "District",
        "Angul",
        "Balasore",
        "Bargarh",
    ]
    assert [region.page for region in table.regions] == [1, 2]
    assert len(table.joined_from) == 2 and table.continues_to is None


def test_a_new_table_with_its_own_caption_above_is_not_linked() -> None:
    doc = _continued_report()
    # The table on page 2 gets its own caption above it ("ESTIMATED LOSS").
    second = doc.tables[1]
    box = BoundingBox(l=50, t=30, r=550, b=50, coord_origin=CoordOrigin.TOPLEFT)
    caption = doc.add_text(
        label="caption",
        text="ESTIMATED LOSS",
        prov=ProvenanceItem(
            page_no=2, bbox=box.to_bottom_left_origin(H), charspan=(0, 14)
        ),
    )
    second.captions.append(caption.get_ref())
    records, _ = extract_tables(doc, _context(), TableExtractionSettings(enabled=True))
    assert records[0].continues_to is None and records[1].continues_from is None


def test_ignored_text_layer_makes_tables_come_from_ocr(service) -> None:
    """Docling's table model reads words from the PDF text layer even in full-page
    OCR mode; with `ignore_text_layer` they must come from OCR instead (PDFs with
    legacy-font text layers)."""
    request = {
        **TABLES_ON,
        "settings": {"ignore_text_layer": True},
        "formats": {
            "pdf": {
                "pipeline_options": {
                    "ocr_options": {
                        "__variant__": "RapidOcrOptions",
                        "mode": "full_page",
                    }
                }
            }
        },
    }
    envelope = _convert(service, TABLE_PAGE.name, TABLE_PAGE.read_bytes(), request)
    assert envelope.status == "success", envelope.errors
    [table] = envelope.tables
    assert table.extraction.method == "ocr"
    assert table.num_cols >= 6
